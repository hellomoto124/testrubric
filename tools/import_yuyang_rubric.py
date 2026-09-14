import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/xiexiaodong/Downloads/雨旸素养框架.xlsx")
OUTPUT = Path("backend/app/data/kai_feng_yuyang_260910.json")


def bullets(value):
    return [item.removeprefix("•").strip() for item in str(value or "").splitlines() if item.strip()]


def parse(source=SOURCE):
    sheet = load_workbook(source, data_only=True)["素养量规表-260910"]
    rows = [row for row in range(6, sheet.max_row + 1) if sheet.cell(row, 5).value]
    goal = capability = None
    metadata = {}
    for row in rows:
        goal = sheet.cell(row, 1).value or goal
        capability = sheet.cell(row, 2).value or capability
        metadata[row] = (goal, capability)
    skills = []
    for index, row in enumerate(rows, 1):
        skills.append({"code": f"YY.S{index:02}", "subskill": sheet.cell(row, 4).value, "name_cn": sheet.cell(row, 5).value, "name_en": sheet.cell(row, 6).value, "source": sheet.cell(row, 7).value, "excel_row": row, "levels": [{"level": level, "descriptors": bullets(sheet.cell(row, 7 + level).value)} for level in range(1, 7)]})
    capabilities = []
    for name in dict.fromkeys(metadata[row][1] for row in rows):
        capability_rows = [skill for skill in skills if metadata[skill["excel_row"]][1] == name]
        capabilities.append({"code": f"YY.C{len(capabilities)+1:02}", "goal": metadata[capability_rows[0]["excel_row"]][0], "name": name, "description": sheet.cell(capability_rows[0]["excel_row"], 3).value, "skills": capability_rows})
    goals = [{"code": f"YY.G{index:02}", "name": name} for index, name in enumerate(dict.fromkeys(metadata[row][0] for row in rows), 1)]
    return {"rubric_version": "kai_feng_yuyang_260910", "school_rule": {"grades": "3-6", "bands": {"basic": 2, "mastery": 3, "excellent": 4}, "default_target": 3}, "goals": goals, "capabilities": capabilities}


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(parse(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
